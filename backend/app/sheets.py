"""Read SAP ID and Band (Mhz) columns from an uploaded Excel / CSV file."""
import csv
from pathlib import Path


class SheetError(ValueError):
    pass


def _clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _rows(path: Path, ext: str):
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
            yield from csv.reader(fh)
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            yield from wb.active.iter_rows(values_only=True)
        finally:
            wb.close()
    elif ext == ".xls":
        import xlrd

        sh = xlrd.open_workbook(path).sheet_by_index(0)
        for i in range(sh.nrows):
            yield sh.row_values(i)
    else:
        raise SheetError("Only Excel (.xls, .xlsx) and CSV files are allowed")


def parse_sheet(path: Path, ext: str) -> tuple[list[str], list[str]]:
    it = iter(_rows(path, ext))
    try:
        header = [_clean(h).lower() for h in next(it)]
    except StopIteration:
        raise SheetError("File is empty")
    sap = next((i for i, h in enumerate(header) if "sap" in h), None)
    band = next((i for i, h in enumerate(header) if "band" in h), None)
    if sap is None or band is None:
        raise SheetError('File needs "SAP ID" and "Band (Mhz)" columns')
    saps, bands = {}, {}
    for row in it:
        if len(row) > max(sap, band):
            s, b = _clean(row[sap]), _clean(row[band])
            if s:
                saps[s] = None
            if b:
                bands[b] = None
    if not saps or not bands:
        raise SheetError("No SAP ID / Band values found in the file")

    def num(x: str):
        try:
            return (0, float(x))
        except ValueError:
            return (1, x)

    return list(saps), sorted(bands, key=num)
